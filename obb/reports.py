"""Reporte de desempeño por estrategia a partir del Strategy Report de Thinkorswim (CSV / XLSX)."""
import csv
import io
import re
import time
import unicodedata
from datetime import datetime

from . import db

EXPECTED = ["Martillo Bajista", "Vela Roja de Apertura", "Vela Verde"]  # se muestran aunque tengan 0 operaciones
DATE_FORMATS = ["%m/%d/%y %I:%M %p", "%m/%d/%Y %I:%M %p", "%m/%d/%y %H:%M", "%m/%d/%Y %H:%M", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M", "%m/%d/%y", "%m/%d/%Y", "%Y-%m-%d"]


class ReportError(ValueError):
    pass


def _norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"\s+", " ", s)


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s in ("-", "--"):
        return None
    neg = s.startswith("(") and s.endswith(")") or s.startswith("-") or "-$" in s or "$-" in s
    s = re.sub(r"[^\d.]", "", s.replace(",", ""))
    if not s or s == ".":
        return None
    n = float(s)
    return -n if neg else n


def _dt(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, (int, float)) and 30000 < v < 80000:  # serie de fecha de Excel
        from datetime import timedelta
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).strftime("%Y-%m-%d %H:%M")
    s = str(v).strip()
    for f in DATE_FORMATS:
        try:
            return datetime.strptime(s, f).strftime("%Y-%m-%d %H:%M")
        except ValueError:
            pass
    return s  # texto desconocido: se conserva tal cual


# ---------------------------------------------------------------- lectura
def _read_rows(filename: str, data: bytes):
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from . import xlsxlite
        sheets = xlsxlite.sheets(data)
        for n, rows in sheets:
            if _norm(n) in ("datos", "data", "trades", "trade list"):
                return rows
        return sheets[0][1]
    if name.endswith(".xls"):
        raise ReportError("El formato .xls antiguo no se admite: guárdalo como .xlsx o CSV.")
    text = None
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeError:
            continue
    if text is None:
        raise ReportError("No se pudo leer el archivo")
    lines = [l for l in text.splitlines() if l.strip()]
    probe = next((l for l in lines if _norm(l).count("strateg") or _norm(l).count("estrateg")), lines[0] if lines else "")
    delim = max([";", ",", "\t"], key=lambda d: probe.count(d))
    return [r for r in csv.reader(lines, delimiter=delim)]


ALIAS = {
    "strategy": ["strategy", "estrategia"],
    "side": ["side", "lado"],
    "when": ["date/time", "datetime", "fecha/hora", "fecha y hora", "fecha", "date", "hora", "time"],
    "entry": ["entrada", "entry", "entry time", "apertura"],
    "exit": ["salida", "exit", "exit time", "cierre"],
    "pnl": ["trade p/l", "p/l trade", "p/l operacion", "p/l neto", "pnl", "p&l", "p/l", "ganancia/perdida", "resultado"],
}


def _map_header(row):
    cols = {}
    cells = [_norm(c) for c in row]
    for key, names in ALIAS.items():
        for nm in names:  # el orden de `names` da prioridad (p. ej. 'trade p/l' antes que 'p/l')
            if nm in cells:
                cols[key] = cells.index(nm)
                break
    return cols


def parse(filename: str, data: bytes) -> dict:
    rows = _read_rows(filename, data)
    hdr_i, cols = None, {}
    for i, r in enumerate(rows[:60]):
        m = _map_header(r)
        if "strategy" in m and "pnl" in m:
            hdr_i, cols = i, m
            break
    if hdr_i is None:
        raise ReportError("No encontré las columnas 'Strategy' y 'Trade P/L' (o 'Estrategia' y 'P/L'). "
                          "Exporta el Strategy Report desde Thinkorswim (Trade List > Export) y vuelve a subirlo.")
    meta_txt = " ".join(" ".join(str(c) for c in r if c is not None) for r in rows[:hdr_i])
    sym = re.search(r"(?:symbol|s[ií]mbolo)\s*[:=]?\s*([A-Za-z.]{1,6})\b", meta_txt, re.I)
    symbol = sym.group(1).upper() if sym else ""

    def cell(r, k):
        return r[cols[k]] if k in cols and cols[k] < len(r) else None

    body = [r for r in rows[hdr_i + 1:] if any(c not in (None, "") for c in r)]
    has_open_rows = "side" in cols and any(_num(cell(r, "pnl")) is None for r in body)
    trades, opened, last_strat, warnings = [], {}, "", []
    for r in body:
        strat = str(cell(r, "strategy") or "").strip() or last_strat
        last_strat = strat or last_strat
        if not strat or _norm(strat) in ("total", "totales"):
            continue
        pnl = _num(cell(r, "pnl"))
        if has_open_rows:  # formato TOS: una fila por orden (entrada / salida)
            ts = _dt(cell(r, "when"))
            if pnl is None:
                opened[strat] = ts
            else:
                trades.append({"strategy": strat, "entry_ts": opened.pop(strat, None), "exit_ts": ts, "pnl": pnl})
        else:  # una fila por operación cerrada
            if pnl is None:
                continue
            exit_ts = _dt(cell(r, "exit")) or _dt(cell(r, "when"))
            trades.append({"strategy": strat, "entry_ts": _dt(cell(r, "entry")), "exit_ts": exit_ts, "pnl": pnl})
    if not trades:
        raise ReportError("El archivo no contiene operaciones cerradas (filas con Trade P/L).")
    if opened:
        warnings.append(f"{len(opened)} estrategia(s) tienen una entrada abierta sin cerrar; no se cuentan.")
    stamps = sorted(t for tr in trades for t in (tr["entry_ts"], tr["exit_ts"]) if t and re.match(r"\d{4}-", t))
    return {"symbol": symbol, "trades": trades, "warnings": warnings,
            "period_start": stamps[0] if stamps else "", "period_end": stamps[-1] if stamps else ""}


# ---------------------------------------------------------------- métricas
def _stats(ts):
    pnls = [t["pnl"] for t in ts]
    wins = [p for p in pnls if p > 0]
    losses = [p for p in pnls if p < 0]
    gw, gl = sum(wins), -sum(losses)
    cum, peak, dd = 0.0, 0.0, 0.0
    for p in pnls:
        cum += p
        peak = max(peak, cum)
        dd = max(dd, peak - cum)
    n = len(pnls)
    return {
        "ops": n, "wins": len(wins), "losses": len(losses),
        "hit_rate": (len(wins) / n * 100) if n else None,
        "net": round(sum(pnls), 2),
        "avg_win": round(gw / len(wins), 2) if wins else None,
        "avg_loss": round(-gl / len(losses), 2) if losses else None,
        "best": round(max(pnls), 2) if pnls else 0.0,
        "worst": round(min(pnls), 2) if pnls else 0.0,
        "profit_factor": round(gw / gl, 2) if gl else None,  # None = sin pérdidas (o sin datos)
        "expectancy": round(sum(pnls) / n, 2) if n else None,
        "max_dd": round(dd, 2),
    }


def summarize(trades: list, expected=EXPECTED) -> dict:
    order = sorted(trades, key=lambda t: (t["exit_ts"] or ""))
    by = {}
    for t in order:
        by.setdefault(t["strategy"], []).append(t)
    for exp in expected:  # estrategias esperadas sin operaciones (como "Martillo Bajista" en el reporte)
        if not any(_norm(exp) in _norm(n) for n in by):
            by[exp] = []
    strategies = []
    for name, ts in sorted(by.items(), key=lambda kv: (-len(kv[1]), kv[0])):
        strategies.append({"name": name, **_stats(ts), "absent": not ts})
    total = _stats(order)
    cum, equity = 0.0, []
    for i, t in enumerate(order, 1):
        cum += t["pnl"]
        equity.append({"n": i, "ts": t["exit_ts"], "cum": round(cum, 2), "strategy": t["strategy"], "pnl": t["pnl"]})
    notes = [f"{s['name']} no muestra operaciones cerradas en este periodo." for s in strategies if s["absent"]]
    return {"strategies": strategies, "total": total, "equity": equity, "notes": notes}


# ---------------------------------------------------------------- base de datos
def save_report(filename, symbol, trades, period_start="", period_end="", source="Thinkorswim Strategy Report", is_demo=False):
    rid = db.x(
        "INSERT INTO reports(filename,symbol,period_start,period_end,source,uploaded_at,is_demo) VALUES(?,?,?,?,?,?,?)",
        (filename, symbol, period_start, period_end, source, int(time.time()), 1 if is_demo else 0),
    )
    db.many("INSERT INTO trades(report_id,strategy,entry_ts,exit_ts,pnl) VALUES(?,?,?,?,?)",
            [(rid, t["strategy"], t["entry_ts"], t["exit_ts"], t["pnl"]) for t in trades])
    return rid


def list_reports():
    return [dict(r) for r in db.q("SELECT * FROM reports ORDER BY uploaded_at DESC, id DESC")]


def load_report(rid: int):
    r = db.q1("SELECT * FROM reports WHERE id=?", (rid,))
    if not r:
        return None
    trades = [dict(t) for t in db.q("SELECT strategy,entry_ts,exit_ts,pnl FROM trades WHERE report_id=? ORDER BY exit_ts,id", (rid,))]
    return {"meta": dict(r), "trades": trades, **summarize(trades)}


def delete_report(rid: int):
    db.x("DELETE FROM trades WHERE report_id=?", (rid,))
    db.x("DELETE FROM reports WHERE id=?", (rid,))


def seed_demo():
    """Reporte de ejemplo (cifras ilustrativas) para que el módulo no aparezca vacío."""
    if db.q1("SELECT 1 FROM reports LIMIT 1"):
        return
    V, R = "Vela Verde (OBB_VVV3)", "Vela Roja de Apertura (OBB_VRV)"
    t = [
        (V, "2026-09-23 07:50", "2026-09-23 09:35", -702.10), (V, "2026-09-23 10:15", "2026-09-23 11:40", 1612.40),
        (R, "2026-09-24 07:50", "2026-09-24 08:25", 2056.00), (V, "2026-09-24 09:50", "2026-09-24 10:30", -655.40),
        (V, "2026-09-25 08:05", "2026-09-25 10:10", 1629.19), (V, "2026-09-25 11:20", "2026-09-25 12:05", -688.25),
        (V, "2026-09-28 09:40", "2026-09-28 10:55", -640.18), (V, "2026-09-28 12:10", "2026-09-28 14:30", 1510.33),
        (V, "2026-09-29 08:15", "2026-09-29 09:20", -671.30), (V, "2026-09-29 10:45", "2026-09-29 13:15", 1576.75),
        (V, "2026-09-30 09:30", "2026-09-30 11:00", -664.90),
    ]
    save_report("DEMO_Strategy_Report_MU.csv", "MU", [{"strategy": a, "entry_ts": b, "exit_ts": c, "pnl": d} for a, b, c, d in t],
                "2026-09-23 07:50", "2026-09-30 16:05", source="Datos de demostración", is_demo=True)


# ---------------------------------------------------------------- exportar a Excel
def to_xlsx(rep: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    meta, S, T = rep["meta"], rep["strategies"], rep["total"]
    fill = lambda hexc: PatternFill("solid", fgColor=hexc)
    white = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D0D7E2")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    MONEY, PCT = '"$"#,##0.00;[Red]("$"#,##0.00)', '0.0%'

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col, w in zip("BCDEFGHIJKL", (36, 13, 10, 10, 11, 14, 15, 15, 14, 14, 13)):
        ws.column_dimensions[col].width = w

    ws.merge_cells("B2:L2")
    ws["B2"] = "OPTIONS BB  |  Reporte de Desempeño por Estrategia"
    ws["B2"].font = Font(color="FFFFFF", bold=True, size=18)
    ws["B2"].fill = fill("0A1628")
    ws["B2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 34
    ws.merge_cells("B3:L3")
    ws["B3"] = (f"Símbolo: {meta.get('symbol') or '—'}  |  Periodo: {meta.get('period_start') or '—'} → {meta.get('period_end') or '—'}"
                f"  |  Fuente: {meta.get('source') or 'Thinkorswim Strategy Report'}")
    ws["B3"].font = Font(color="D6E4FF", size=10)
    ws["B3"].fill = fill("0A1628")
    ws["B3"].alignment = Alignment(vertical="center", indent=1)

    kpis = [("B", "B", "OPERACIONES CERRADAS", T["ops"], "0F2A4A", "0"), ("C", "D", "GANADAS", T["wins"], "1B7F3B", "0"),
            ("E", "F", "PERDIDAS", T["losses"], "8B1A1F", "0"), ("G", "H", "% ACIERTO GLOBAL", (T["hit_rate"] or 0) / 100, "0F5F7F", PCT),
            ("I", "L", "P/L NETO TOTAL", T["net"], "0F2A4A", MONEY)]
    for c1, c2, label, val, color, fmt in kpis:
        if c1 != c2:
            ws.merge_cells(f"{c1}5:{c2}5")
            ws.merge_cells(f"{c1}6:{c2}6")
        ws[f"{c1}5"], ws[f"{c1}6"] = label, val
        ws[f"{c1}5"].font, ws[f"{c1}6"].font = Font(color="FFFFFF", bold=True, size=9), Font(color="FFFFFF", bold=True, size=20)
        ws[f"{c1}5"].number_format = ws[f"{c1}6"].number_format = fmt
        for r in (5, 6):
            for col in range(ord(c1), ord(c2) + 1):
                ws[f"{chr(col)}{r}"].fill = fill(color)
            ws[f"{c1}{r}"].alignment = center
    ws.row_dimensions[6].height = 36

    heads = ["Estrategia", "Operaciones", "Ganadas", "Perdidas", "% Acierto", "P/L Neto", "Ganancia Prom.", "Pérdida Prom.",
             "Mejor Op.", "Peor Op.", "Factor Ben."]
    hr = 8
    for j, h in enumerate(heads):
        c = ws.cell(hr, 2 + j, h)
        c.font, c.fill, c.alignment, c.border = white, fill("000000"), center, box
    ws.row_dimensions[hr].height = 28
    r = hr
    for s in S:
        r += 1
        pf = s["profit_factor"]
        vals = [s["name"], s["ops"], s["wins"], s["losses"], (s["hit_rate"] / 100) if s["hit_rate"] is not None else None,
                s["net"], s["avg_win"], s["avg_loss"], s["best"], s["worst"] if s["ops"] else None,
                None if pf in (None, float("inf")) else pf]
        for j, v in enumerate(vals):
            c = ws.cell(r, 2 + j, v)
            c.border, c.alignment = box, Alignment(horizontal="left" if j == 0 else "center")
            c.fill = fill("EEF3FA")
            c.number_format = {4: PCT, 5: MONEY, 6: MONEY, 7: MONEY, 8: MONEY, 9: MONEY, 10: "0.00"}.get(j, "General")
        ws.cell(r, 2).font = Font(color="0F5F7F", bold=True)
    first_s, last_s = hr + 1, r
    r += 1
    tot = ["TOTAL", T["ops"], T["wins"], T["losses"], (T["hit_rate"] or 0) / 100, T["net"], None, None, None, None, None]
    for j, v in enumerate(tot):
        c = ws.cell(r, 2 + j, v)
        c.font, c.fill, c.border, c.alignment = white, fill("000000"), box, Alignment(horizontal="left" if j == 0 else "center")
        c.number_format = {4: PCT, 5: MONEY}.get(j, "General")
    r += 2
    for note in rep["notes"]:
        ws.cell(r, 2, "Nota: " + note).font = Font(italic=True, color="666666", size=9)
        r += 1

    ch1 = BarChart()
    ch1.type, ch1.title, ch1.legend = "col", "% de Acierto por Estrategia", None
    ch1.add_data(Reference(ws, min_col=6, min_row=hr, max_row=last_s), titles_from_data=True)
    ch1.set_categories(Reference(ws, min_col=2, min_row=first_s, max_row=last_s))
    ch1.y_axis.number_format, ch1.y_axis.scaling.min, ch1.y_axis.scaling.max = "0%", 0, 1
    ch1.dataLabels = DataLabelList()
    ch1.dataLabels.showVal = True
    ch1.series[0].graphicalProperties.solidFill = "1180A8"
    ch1.height, ch1.width = 8.5, 15
    ws.add_chart(ch1, f"B{r + 1}")
    ch2 = BarChart()
    ch2.type, ch2.title = "col", "Ganadas vs Perdidas por Estrategia"
    ch2.add_data(Reference(ws, min_col=4, max_col=5, min_row=hr, max_row=last_s), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=2, min_row=first_s, max_row=last_s))
    ch2.series[0].graphicalProperties.solidFill = "1FA34A"
    ch2.series[1].graphicalProperties.solidFill = "DC2434"
    ch2.height, ch2.width = 8.5, 15
    ws.add_chart(ch2, f"G{r + 1}")

    wd = wb.create_sheet("Datos")
    for j, h in enumerate(["#", "Estrategia", "Entrada", "Salida", "P/L", "Acumulado", "Resultado"], 1):
        c = wd.cell(1, j, h)
        c.font, c.fill, c.alignment = white, fill("000000"), center
    cum = 0.0
    for i, t in enumerate(rep["trades"], 1):
        cum += t["pnl"]
        row = [i, t["strategy"], t["entry_ts"], t["exit_ts"], t["pnl"], round(cum, 2), "Ganada" if t["pnl"] > 0 else "Perdida" if t["pnl"] < 0 else "Neutra"]
        for j, v in enumerate(row, 1):
            c = wd.cell(i + 1, j, v)
            if j in (5, 6):
                c.number_format = MONEY
    for j, w in enumerate((5, 34, 18, 18, 14, 14, 11), 1):
        wd.column_dimensions[get_column_letter(j)].width = w
    wd.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
